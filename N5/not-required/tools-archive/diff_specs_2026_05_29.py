# -*- coding: utf-8 -*-
"""One-shot spec-drift audit helper (2026-05-29).

Goal: find content present in OLDER specs but MISSING from the CURRENT
implementation spec, so we can then check whether that missing info is
actually implemented in the live app.

Older specs:
  - JLPT N5 — Consolidated Spec.docx
  - JLPT N5 Grammar Tutor – Functional Spec.docx
  - JLPT-N5-Functional-Spec-v3.1-supplement.md
Current spec:
  - JLPT-N5-Current-Implementation-Spec.md

Dumps full text of each .docx to tools/_spec_extract/<name>.txt for detailed
reading, and prints a heading inventory per file. Read-only on the specs.
Archive to not-required/ after the investigation."""

import re
import sys
import zipfile
from pathlib import Path

# Windows console defaults to cp932; spec titles contain en/em-dashes.
try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass

SPECDIR = Path(__file__).resolve().parent.parent / 'specifications'
OUT = Path(__file__).resolve().parent / '_spec_extract'
OUT.mkdir(exist_ok=True)


def docx_paragraphs(path):
    """Return list of (style_name, text). Falls back to raw XML if needed."""
    try:
        import docx
        d = docx.Document(str(path))
        out = []
        for p in d.paragraphs:
            style = (p.style.name if p.style else '') or ''
            out.append((style, p.text))
        for ti, t in enumerate(d.tables):
            out.append(('__TABLE__', f'[table {ti} start: {len(t.rows)} rows x {len(t.columns)} cols]'))
            for row in t.rows:
                cells = [c.text.replace('\n', ' ').strip() for c in row.cells]
                out.append(('TableRow', ' | '.join(cells)))
        return out
    except Exception as e:
        with zipfile.ZipFile(str(path)) as z:
            xml = z.read('word/document.xml').decode('utf-8', 'replace')
        xml = re.sub(r'</w:p>', '\n', xml)
        txt = re.sub(r'<[^>]+>', '', xml)
        return [('?', ln) for ln in txt.split('\n')]


def dump_docx(name):
    path = SPECDIR / name
    paras = docx_paragraphs(path)
    safe = re.sub(r'[^A-Za-z0-9._-]+', '_', name)
    outpath = OUT / (safe + '.txt')
    with open(outpath, 'w', encoding='utf-8') as f:
        for style, txt in paras:
            tag = f'[{style}] ' if style and style != 'Normal' else ''
            f.write(f'{tag}{txt}\n')
    # heading inventory
    headings = [(s, t) for s, t in paras
                if ('Heading' in s or 'Title' in s) and t.strip()]
    print(f'\n===== {name} =====')
    print(f'  paragraphs={len(paras)}  headings={len(headings)}  -> {outpath.name}')
    for s, t in headings:
        lvl = ''.join(ch for ch in s if ch.isdigit())
        indent = '  ' * (int(lvl) if lvl else 1)
        print(f'    {indent}- ({s}) {t.strip()[:90]}')


def md_headings(name):
    path = SPECDIR / name
    print(f'\n===== {name} =====')
    txt = path.read_text(encoding='utf-8', errors='replace')
    n = 0
    for ln in txt.splitlines():
        m = re.match(r'^(#{1,4})\s+(.*)', ln)
        if m:
            n += 1
            depth = len(m.group(1))
            print(f'    {"  "*(depth-1)}- {m.group(2).strip()[:90]}')
    print(f'  ({n} headings)')


def list_docx():
    return sorted(p.name for p in SPECDIR.glob('*.docx'))


if __name__ == '__main__':
    print('DOCX files found:', list_docx())
    for name in list_docx():
        dump_docx(name)
    md_headings('JLPT-N5-Current-Implementation-Spec.md')
    md_headings('JLPT-N5-Functional-Spec-v3.1-supplement.md')
    print('\nDone. Full docx text dumped to', OUT)
