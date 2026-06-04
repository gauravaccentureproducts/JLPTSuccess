import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

doc = Document('docs/vocab-review/N5-vocab-995-entries-for-native-review.docx')
full = '\n'.join(p.text for p in doc.paragraphs)
# also tables
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                full += '\n' + p.text

print('Docx すむ-fix verification:')
checks = {
    '東京に すむ (new, should be PRESENT)': '東京に すむ',
    '日本に すむ (new, should be PRESENT)': '日本に すむ',
    'この町に すむ (new, should be PRESENT)': 'この町に すむ',
    'まいにち すむ (old bug, should be ABSENT)': 'まいにち すむ',
    'いま すむ (old bug, should be ABSENT)': 'いま すむ',
    'よく すむ (old, should be ABSENT)': 'よく すむ',
}
for label, needle in checks.items():
    present = needle in full
    print(f'  [{"PRESENT" if present else "absent"}] {label}')

# residual: any form+masu malformed conjugation in the whole docx? (spot the すむ ones)
print()
for bad in ('すむます', 'すむました'):
    print(f'  malformed {bad!r} in docx: {bad in full}')
print(f'\nTotal tables (reviewer-notes boxes): {len(doc.tables)}')
