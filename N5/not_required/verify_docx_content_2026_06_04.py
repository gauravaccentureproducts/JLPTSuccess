"""Check that the regenerated vocab review docx actually contains the
batch A-G content markers. Read-only sanity check."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

doc = Document('docs/vocab-review/N5-vocab-995-entries-for-native-review.docx')
print(f'Total paragraphs: {len(doc.paragraphs)}')
print(f'Total tables:     {len(doc.tables)}')
print()

# Each tuple: (label, content marker string we EXPECT to find verbatim)
markers = [
    ('A   kata fix',              "あの方は"),
    ('B-x ame/雨 false_friends',   "雨"),
    ('B-s same-kana homonym',     "Same-kana homonym"),
    ('D   uchi/soto family',      "Use 母 when referring to YOUR own"),
    ('F1  legacy markers',        "use 電車 instead in modern Japanese"),
    ('F2  naku split',            "鳴く is for animal sounds"),
    ('F3  kanojo caution',        "In everyday modern Japanese, かのじょ"),
    ('F4  sensei scope',          "respectful title used for any expert"),
    ('F5  shitsurei before',      "BEFORE doing something potentially intrusive"),
    ('F5  sayounara formal',      "surprisingly formal in modern usage"),
    ('F6  umi/mizuumi',           "lake (use みずうみ for lake)"),
    ('F6  gohan/kome',            "Cooked rice"),
    ('F7  suki na-adj',           "but すき is a な-adjective"),
    ('F8  kirei dual',            "pretty/beautiful"),
    ('G1  student scope',         "STUDENTS in higher education"),
    ('G2  police register',       "formal / written / news-register word"),
    ('G3  gaikokujin caution',    "literally means 'foreign-country person'"),
    ('G4  ichinichi reading',     "ついたち (1st day of month)"),
    ('G5  konya/konban',          "slightly more formal, written, or newscaster"),
    ('G7  ippai dual',            "'full / brimming'"),
    ('G8  tada multi-sense',      "has several meanings"),
    ('G9  shika + neg',           "REQUIRES a negative predicate"),
    ('G10 gurai/kurai',           "free variants"),
    ('G11 yahari formality',      "formal/written register"),
    ('G12 kotoba/tango',          "covers a wide range: a single word"),
    ('G13 nande sense',           "by what means / by what method"),
    ('G14 hoshii adj',            "い-adjective, not a verb"),
    ('G15 gozaru',                "almost never used in its plain dictionary form"),
    ('G16 hayai split',           "早い means 'early'"),
    ('G17 yasashii split',        "易しい = 'easy / simple'"),
    ('G18 toilet politeness',     "polite Japanese-style word for restroom"),
    ('G19 ofuro/shower',          "TUB-style bath"),
    ('G20 se back/height',        "covers both 'back' (the body part) and 'height"),
    ('G20 ashi leg/foot',         "covers the entire lower limb"),
    ('G21 ji generic',            "broadest term for a written character"),
    ('G21 kana umbrella',         "umbrella term for the two Japanese syllabic systems"),
    ('G21 katakana',              "loanwords"),
    ('G22 karai',                 "primarily means 'spicy / hot'"),
    ('G24 color noun',            "NOUN form of the color"),
    ('G26 hairbrush',             "歯ブラシ (kanji 歯"),
]

para_text = '\n'.join(p.text for p in doc.paragraphs)
hits = 0
misses = []
for label, needle in markers:
    if needle in para_text:
        hits += 1
        print(f'  [OK]   {label}')
    else:
        misses.append((label, needle))
        print(f'  [MISS] {label}')

print()
print(f'Hits: {hits}/{len(markers)}')
if misses:
    print()
    print('--- Misses (needles not found in docx) ---')
    for label, needle in misses:
        print(f'  {label}: {needle!r}')

# Also verify the counter cleanup landed (家族 / 両親 no longer carry counter line)
print()
print('=== Counter cleanup sanity check (家族/両親 should not have ~人 counter) ===')
for needle in ('Counter: 〜にん', 'Counter: ~にん'):
    if needle in para_text:
        # Check whether it appears in the 家族 or 両親 entry's page block
        # Just report raw presence/absence; learner can grep
        print(f'  Found in document: {needle!r} (may be other entries; verify 家族/両親 pages by hand)')
