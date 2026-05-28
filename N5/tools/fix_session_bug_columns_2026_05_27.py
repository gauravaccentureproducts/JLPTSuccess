# -*- coding: utf-8 -*-
"""Swap cols 9-10 on the 5 session bugs I added (BUG-223..227, rows
226-230). The append script put commit-hash in col 9 and fix-date in
col 10, but the actual convention (per JA-146 check at line 8237 of
check_content_integrity.py and per rows 220+ data) is:

  col 9:  fix date (YYYY-MM-DD)
  col 10: commit hash (or '<...>'-shaped sentinel)

The header labels at row 3 ("Fix Commit" at col 9, "Fix Date" at col
10) are misleading — actual data has been in col 10 for hash since
the table was authored. JA-146 reads col 10 as fix-commit.

One-shot script; archive to not-required/ after run."""

import openpyxl
from pathlib import Path

XLSX = Path(__file__).resolve().parent.parent / 'specifications' / 'test-scenarios-by-specialist-perspective.xlsx'


def main():
    wb = openpyxl.load_workbook(XLSX)
    ws = wb['User Reported Bugs']
    # Rows 226-230 are the 5 BUG-223..227 entries I added
    for r in range(226, 231):
        col9 = ws.cell(row=r, column=9).value
        col10 = ws.cell(row=r, column=10).value
        print(f'  before r{r}: col9={col9!r}  col10={col10!r}')
        # Swap
        ws.cell(row=r, column=9, value=col10)
        ws.cell(row=r, column=10, value=col9)
        print(f'  after  r{r}: col9={col10!r}  col10={col9!r}')
    wb.save(XLSX)
    print('Saved.')


if __name__ == '__main__':
    main()
