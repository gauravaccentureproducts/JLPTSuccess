"""Register two BUG rows: Batch H (Fixed) + roll-up of remaining OPEN items
from the 2026-06-04 follow-up review."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from copy import copy
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / 'specifications' / 'test-scenarios-by-specialist-perspective.xlsx'
wb = load_workbook(XLSX)
ws = wb['User Reported Bugs']
src_row = ws.max_row


def insert_row(values):
    new_row = ws.max_row + 1
    for col in range(1, ws.max_column + 1):
        src_cell = ws.cell(src_row, col)
        new_cell = ws.cell(new_row, col)
        if col in values:
            new_cell.value = values[col]
        if src_cell.has_style:
            new_cell.font = copy(src_cell.font)
            new_cell.alignment = copy(src_cell.alignment)
            if src_cell.fill.fill_type:
                new_cell.fill = copy(src_cell.fill)
    return new_row


# Row 1 — Batch H Fixed
r_h = insert_row({
    1: '="BUG-"&TEXT(ROW()-3,"000")',
    2: '2026-06-04',
    3: 'Native Japanese / JLPT expert reviewer (follow-up consolidated report)',
    4: 'Native-review batch H (23 entries) — mechanical closure of remaining OPEN items from consolidated follow-up review',
    5: (
        '23 entries touched + corpus-level review-status disclosure added. '
        'H1: removed counter from りょこう (occurrence-counting via verb, '
        'not direct noun counter). H2: added "legacy / low-frequency in '
        'modern daily Japanese" pragmatic notes to レコード, フィルム, マッチ, '
        'はいざら (counters retained — they are correct for the noun shape). '
        'H3: verified existing transitivity + pair_id labels for 6 verb '
        'pairs (あく/あける, しまる/しめる, おちる/おとす, でる/だす, きえる/'
        'けす, はじまる/はじめる) — all already correctly populated, no '
        'action needed. H4: pronoun cautions added to 私, 私たち, みんな, '
        'みなさん (Japanese subject omission preference + みんな casual '
        'vs みなさん polite contrast). H5: 14 family entries received '
        'the native reviewer\'s preferred concise JA wording as '
        'supplementary pragmatic notes (complements the longer English '
        'explanations from batch D). H6: added _meta.'
        'native_human_review_status_2026_06_04 disclosure explicitly '
        'stating the corpus is AI-reviewed with native-human review '
        'pending (addresses OPEN-002 concern at corpus level).'
    ),
    6: 'Medium',
    7: 'P2',
    8: 'Fixed',
    9: '2026-06-04',
    10: '<pending-commit-hash>',
    11: '',
    12: '',
    13: 'BUG-262 (A-D), BUG-263 (E open), BUG-264 (F), BUG-265 (G)',
    14: 'tools/apply_native_review_batch_h_2026_06_04.py + CI 179/179',
})

# Row 2 — remaining-OPEN-items roll-up
r_open = insert_row({
    1: '="BUG-"&TEXT(ROW()-3,"000")',
    2: '2026-06-04',
    3: 'Native Japanese / JLPT expert reviewer (follow-up consolidated report)',
    4: 'OPEN follow-up: remaining native-reviewer-pending items that cannot be auto-closed',
    5: (
        'Follow-up consolidated review identified 13 OPEN-NNN classes. '
        'Auto-closure addressed OPEN-006 (homonyms partial), OPEN-008 '
        '(pronouns partial via H4), OPEN-009 (family terms via D + H5), '
        'OPEN-010 (counter cleanup via C + H1), OPEN-011 (legacy '
        'markers via F1 + H2), OPEN-012 (verb-pair labels — already '
        'populated in corpus). REMAINING OPEN — these require an '
        'actual human native speaker, NOT Claude impersonating one: '
        'OPEN-001 every-entry 「問題なし」 / 「要修正」 verdict (995 '
        'entries; cannot be auto-written without dishonesty); '
        'OPEN-002 entry-level "ai_reviewed_only" vs '
        '"native_reviewed" status discipline at the data-level '
        '(corpus-level disclosure shipped via H6; per-entry flag '
        'would need a schema field add); OPEN-003 corpus-wide '
        'particle-example rewrites (already tracked as BUG-263 '
        'Open); OPEN-004 template-collocation cleanup (same root '
        'cause as OPEN-003; tracked under BUG-263); OPEN-005 '
        'English-gloss reframing (case-by-case native judgment); '
        'OPEN-007 kana-only display ambiguity (requires UI'
        '-level renderer work to surface kanji disambiguation '
        'on kana-only entries — different cycle, may be a JA-'
        'invariant or kanji-tooltip work item); OPEN-013 example-'
        'sentence quality pass (case-by-case native judgment). '
        'These remain in the queue for a real human native-reviewer '
        'closure cycle.'
    ),
    6: 'Medium',
    7: 'P2',
    8: 'Open',
    9: '',
    10: '<awaiting-native-pass>',
    11: '',
    12: '',
    13: 'BUG-263 (batch E open), BUG-262/264/265/266 (batches A-H done)',
    14: 'Awaiting human native-reviewer per-entry pass',
})

wb.save(XLSX)
print(f'Inserted Batch H (row {r_h}) + OPEN-followup roll-up (row {r_open}).')
