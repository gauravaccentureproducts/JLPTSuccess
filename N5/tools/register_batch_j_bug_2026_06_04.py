"""Register BUG row for native-review Batch J (systematic template artifacts)."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from copy import copy
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / 'specifications' / 'test-scenarios-by-specialist-perspective.xlsx'
wb = load_workbook(XLSX)
ws = wb['User Reported Bugs']
src_row = ws.max_row
new_row = ws.max_row + 1

values = {
    1: '="BUG-"&TEXT(ROW()-3,"000")',
    2: '2026-06-04',
    3: 'Native Japanese / JLPT expert reviewer (specific-bug audit)',
    4: 'Native-review batch J — systematic template-artifact fixes (142 entries): malformed verb conjugations + nonsense template example sentences',
    5: (
        'A specific-bug audit cited 12 concrete JP-content errors. '
        'Reproduction confirmed all 12 AND revealed they are instances of '
        'systematic template-generation artifacts. Batch J fixed the '
        'deterministic + confirmed-nonsense classes (142 entries): '
        'J1 — 110 verbs had malformed 「<dict-form>ます」/「<dict-form>ました」 '
        'in particle_examples (you cannot append ます to a dictionary '
        'form). Corrected to proper polite/polite-past via a verb-class-'
        'aware conjugator (68 godan, 31 ichidan, 11 irregular). E.g. '
        'べんきょうするます→べんきょうします, かすました→かしました, あるます→あります, '
        'しごとするました→しごとしました, すむます→すみます. 0 malformed remain. '
        'J2 — removed 16 「えいがは <i-adj>でした。」 example sentences (both '
        'grammatically wrong — i-adj past is かったです not でした — AND '
        'semantically nonsense, "the movie was yellow/thin/bitter") + 1 '
        'stray cross-contaminated particle_example (しかくいでした in 書く). '
        'J3 — removed 14 「今日は とても <X>です。」 example sentences for '
        'non-weather adjectives (今日は 白い/くろい/みじかい/まるい/ぬるい/とおい/'
        'おおい… = nonsense; kept weather adjectives すずしい/あたたかい/いい). '
        'J4 — cited specifics: cleaned カタカナ particle_examples (removed '
        'purchase/price/new template collocations カタカナを かう/を ください/'
        'あたらしい/たかい/やすい — a writing system is not a buyable object); '
        'removed 「この りんごは うるさいです。」 (an apple cannot be noisy). '
        'Of the 12 cited bugs, 10 are fully fixed here; the remaining 2 '
        '(まいにち すむ / いま すむ adverb-collocations in particle_examples) '
        'are part of the broad OPEN-003/004 particle_examples-naturalness '
        'class tracked under BUG-263 — they are grammatical but '
        'stylistically weak, and fixing them well needs the same per-verb '
        'native judgment as the other ~109 verbs adverb collocations. '
        'Driver: tools/apply_native_review_batch_j_2026_06_04.py. '
        'CI: 179/179 PASS.'
    ),
    6: 'High',
    7: 'P1',
    8: 'Fixed',
    9: '2026-06-04',
    10: '<pending-commit-hash>',
    11: '',
    12: '',
    13: 'BUG-263 (OPEN-003/004 particle-example naturalness — remaining 2 cited + broad template class)',
    14: 'tools/apply_native_review_batch_j_2026_06_04.py + verify_batch_j (0 malformed / 0 nonsense remaining) + CI 179/179',
}

for col in range(1, ws.max_column + 1):
    src_cell = ws.cell(src_row, col)
    new_cell = ws.cell(new_row, col)
    if col in values:
        new_cell.value = values[col]
    if src_cell.has_style:
        new_cell.font = copy(src_cell.font)
        new_cell.alignment = copy(src_cell.alignment)
        if src_cell.fill.fill_type:
            new_cell.fill = copy(src_cell.fill)

wb.save(XLSX)
print(f'Inserted Batch J row {new_row}')
