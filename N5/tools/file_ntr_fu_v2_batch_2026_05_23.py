"""File NTR-FU-004 / 005 / 006 / 007 in the bug tracker.

Four bugs from the 2026-05-23 re-review against v1.16.2.
Verified-real per F.41.4 + the F.44.17 amendment (actual-data-
inspection mandatory, not commit-message-trust). My earlier
classification of NTR-FU-004's parent claim as STALE was a
verification-script bug (wrong top-level key in grammar.json
loader); the deprecation flags ARE set, but the follow-up
discipline (core_n5 cleanup + contrasts.note staleness) remained.
"""
import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
import openpyxl

REPO_N5 = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(REPO_N5, "specifications", "test-scenarios-by-specialist-perspective.xlsx")

BUGS = [
    {
        "row": 180,
        "title": "NTR-FU-004 — n5-045 deprecated but still in core_n5 list + contrasts.note still self-identifies as duplicate",
        "description": (
            "**Source:** 2026-05-23 re-review of v1.16.2 (review packet v1.16.2 build).\n\n"
            "**Issue:** NTR-002 close (eda9441) correctly set "
            "n5-045.deprecated=True + n5-045._alias_of='n5-017' on the grammar.json side, "
            "and cleared the reverse _alias_of on n5-017. But two pieces of follow-up "
            "discipline remained:\n\n"
            "  - **A) n5_core_pattern_ids.json:** n5-045 is still in `core_n5` list "
            "(153 entries). Deprecated entries shouldn't appear in the canonical N5 "
            "pattern catalog — they should be moved to a new `deprecated` bucket so "
            "consumers can filter consistently.\n\n"
            "  - **B) n5-045.contrasts[0].note:** still reads 'This is a duplicate "
            "entry - see the canonical pattern.' Now redundant with the new "
            "deprecation lattice (deprecated=True + _alias_of + deprecated_reason). "
            "Update to point at the alias structure, or remove.\n\n"
            "**Bonus discipline lesson:** my prior re-paste triage (Part 42) classified "
            "this item as STALE based on commit-message-trust + a buggy verification "
            "script (used `g.get('grammar')` instead of `g.get('patterns')`); never "
            "actually loaded data. F.44.17 needs an amendment: STALE classification "
            "MUST run actual data-inspection that prints non-empty per-claim output.\n\n"
            "**Fix:** move n5-045 from `core_n5` to new `deprecated` bucket; update "
            "n5-045.contrasts[0].note; add JA-153 CI invariant locking the cleanup."
        ),
        "severity": "Medium",
        "priority": "P3",
    },
    {
        "row": 181,
        "title": "NTR-FU-005 — あなた examples [1] and [2] still use あなた the way the new usage_note tells learners not to",
        "description": (
            "**Source:** 2026-05-23 re-review, item #6 follow-up.\n\n"
            "**Issue:** NTR-004 close (eda9441) added a usage_note to あなた "
            "documenting the formal/intimate/marked-only restriction + replaced "
            "example [0] with a name+さん alternative ('田中さんは どこから 来ましたか。'). "
            "Examples [1] and [2] were NOT touched in that pass and contradict the "
            "new usage_note:\n\n"
            "  - [1] 'あなたは がくせいですか。' — generic 'are you a student?' uses "
            "あなた the way the usage_note warns against\n"
            "  - [2] 'あなたは 何さいですか。' — same pattern, generic 'how old are you?'\n\n"
            "**Fix:** rewrite [1] to parallel [0] (use name+さん); rewrite [2] to "
            "demonstrate the actual appropriate context (form-filling: 'あなたの 名前を "
            "ここに 書いて ください。'). Preserves one あなた example showing where "
            "it IS the right pronoun choice."
        ),
        "severity": "Medium",
        "priority": "P3",
    },
    {
        "row": 182,
        "title": "NTR-FU-006 — じぶん (reflexive pronoun) carries plain counter='人/にん' despite being reflexive",
        "description": (
            "**Source:** 2026-05-23 re-review, item #12 follow-up.\n\n"
            "**Issue:** NTR-013 close added applies_to='noun_of_reference' annotation "
            "to collective pronouns (私たち, みなさん). The singular pronouns (私, あなた, "
            "かれ, かのじょ, かた, 人) keep plain counter — defensible because counting "
            "people with 人 (一人, 二人, 三人) IS semantically valid for singulars.\n\n"
            "But じぶん is a **reflexive** pronoun, not a singular pronoun in the same "
            "sense. 'How many ones-self' isn't a meaningful count — applies_to='noun_of_"
            "reference' is the more honest annotation. NTR-013's bounded close was "
            "correct for collectives; じぶん is an outlier in the singular bucket.\n\n"
            "**Fix:** add applies_to='noun_of_reference' + note explaining reflexive "
            "doesn't count itself."
        ),
        "severity": "Low",
        "priority": "P4",
    },
    {
        "row": 183,
        "title": "NTR-FU-007 — おはし + えいが ID slugs reference legacy section numbers (20-tableware + 26-house) after section field was retagged",
        "description": (
            "**Source:** 2026-05-23 re-review, new finding (not in original NTR batch).\n\n"
            "**Issue:** NTR-005 (BUG-165) + NTR-006 (BUG-166) corrected section field:\n"
            "  - おはし: section field → '19. Tableware and Cooking' (correct)\n"
            "  - えいが: section field → '37. Common nouns - miscellaneous' (correct)\n\n"
            "But the entry IDs still embed the OLD section numbers:\n"
            "  - おはし id: 'n5.vocab.20-tableware-and-cooking.はし-chopsticks'\n"
            "  - えいが id: 'n5.vocab.26-house-and-furniture.えいが'\n\n"
            "IDs are referenced by audio_manifest, questions.json, user localStorage, "
            "etc. Renaming would break references. The right call is **keep IDs "
            "immutable + treat section FIELD as authoritative** + flag the legacy "
            "embedding so future audits / consumers don't parse ID for section.\n\n"
            "**Fix:** add `legacy_section_in_id: true` flag to both entries + add "
            "JA-154 CI invariant warning when any vocab entry's ID-slug-encoded "
            "section disagrees with its `section` field (without the legacy flag)."
        ),
        "severity": "Low",
        "priority": "P4",
    },
]


def main():
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["User Reported Bugs"]
    for b in BUGS:
        r = b["row"]
        ws.cell(r, 1).value = '="BUG-"&TEXT(ROW()-3,"000")'
        ws.cell(r, 2).value = "2026-05-23"
        ws.cell(r, 3).value = "native-teacher review re-pass against v1.16.2 2026-05-23"
        ws.cell(r, 4).value = b["title"]
        ws.cell(r, 5).value = b["description"]
        ws.cell(r, 6).value = b["severity"]
        ws.cell(r, 7).value = b["priority"]
        ws.cell(r, 8).value = "Open"
        ws.cell(r, 9).value = None
        ws.cell(r, 10).value = None
        print(f"  Filed row {r}: BUG-{r-3:03d} — {b['title'][:60]}")
    wb.save(XLSX)
    print()
    print(f"=== Filed {len(BUGS)} new bugs (BUG-177..180) ===")


if __name__ == "__main__":
    main()
