"""
Build CategorizedtestScenarios.xlsx — parallel user-facing view of test scenarios.

Source : data/grammar.json (178 patterns)
Output : test/categorized testing/CategorizedtestScenarios.xlsx

Sheets:
  1. ReadMe
  2. Grammar                  - 10 test scenarios applied to each pattern
  3. Grammar Pattern List     - 178 rows, one per pattern, 10 result columns
  4. Moji                     - scaffold only (Phase 2)
  5. Goi                      - scaffold only (Phase 2)
  6. Dokkai                   - scaffold only (Phase 2)
  7. Chokai                   - scaffold only (Phase 2)
  8. Test Papers              - scaffold only (Phase 2)
"""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


# ---------- paths ----------

ROOT = Path(__file__).resolve().parent.parent
GRAMMAR_JSON = ROOT / "data" / "grammar.json"
OUT_DIR = ROOT / "test" / "categorized testing"
OUT_FILE = OUT_DIR / "CategorizedtestScenarios.xlsx"

SPA_URL_TEMPLATE = "https://gauravaccentureproducts.github.io/JLPTSuccess/N5/#/learn/{id}"


# ---------- styling ----------

GREEN = "14452A"
GREEN_SOFT = "1A5C38"
HEADER_FILL = PatternFill("solid", fgColor=GREEN)
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
SUBHEADER_FILL = PatternFill("solid", fgColor="D9E5DD")
SUBHEADER_FONT = Font(bold=True, color=GREEN, size=10)
WRAP = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER = Alignment(horizontal="center", vertical="center")
THIN = Side(style="thin", color="C0C0C0")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)


def style_header_row(ws, row=1):
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def auto_width(ws, widths):
    for col_idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = w


def freeze_top(ws):
    ws.freeze_panes = "A2"


# ---------- scenario definitions (10 for Grammar) ----------

GRAMMAR_SCENARIOS = [
    # (ts_id, name, description, category, test_steps, expected_result)
    ("TS-01", "JA grammar — Explanation paragraph",
     "Check the Japanese-language correctness of the explanation paragraph for the pattern.",
     "Japanese language correctness (6.1)",
     "1. Open the pattern's interactive page (URL in Grammar Pattern List sheet). "
     "2. Read the Japanese explanation section (when authored — Phase 2). "
     "3. Verify: hiragana/katakana/kanji spelling, particle usage, verb conjugation, tense, politeness register.",
     "No grammar/spelling errors; politeness register matches pattern; particles correct."),

    ("TS-02", "JA grammar — Example sentences",
     "Check the Japanese-language correctness of every example sentence for the pattern.",
     "Japanese language correctness (6.1)",
     "1. Open the pattern's interactive page. "
     "2. For each of ~10 example sentences, verify the JA grammar, particle usage, and conjugation. "
     "3. Verify the sentence actually uses the pattern being taught.",
     "All example sentences are grammatically correct AND demonstrate the pattern."),

    ("TS-03", "JA grammar — Common mistakes (wrong/right pairs)",
     "Check the Japanese in both the 'wrong' and 'corrected' lines is internally consistent and the wrong example is a realistic learner mistake.",
     "Japanese language correctness (6.1)",
     "1. Open the pattern's interactive page. "
     "2. For each wrong/right pair: verify the WRONG line is actually wrong (realistic learner mistake) and the RIGHT line is genuinely correct.",
     "Wrong/right pairs are accurate; wrong line is a plausible learner error; right line uses the target pattern correctly."),

    ("TS-04", "Consistency — Pattern name matches explanation",
     "Check that the pattern name (header), the meaning line, and the body of the explanation describe the same grammar point.",
     "Pattern/explanation/example consistency (6.2)",
     "1. Read the pattern name in the header. "
     "2. Read the meaning line (one-liner). "
     "3. Read the full explanation. "
     "4. Verify all three describe the same grammar point with no contradiction.",
     "Pattern name, meaning line, and explanation describe the SAME pattern with no semantic drift."),

    ("TS-05", "Consistency — Examples demonstrate the pattern",
     "Every example sentence must actually USE the pattern being taught (not just be tangentially related).",
     "Pattern/explanation/example consistency (6.2)",
     "1. Identify the target pattern (e.g., 〜です／〜ます). "
     "2. For each example sentence, verify the pattern is actually present and used as the central grammar point.",
     "≥90% of examples directly demonstrate the pattern. Any contrast examples are clearly labelled."),

    ("TS-06", "Consistency — Audio matches written JA",
     "The audio file for each example must play and audibly match the written Japanese sentence.",
     "Pattern/explanation/example/audio consistency (6.2)",
     "1. Click the audio button for each example. "
     "2. Listen and verify the audio matches the displayed Japanese sentence (correct words, correct order, no truncation, no extra content).",
     "Audio plays cleanly AND matches the displayed text word-for-word."),

    ("TS-07", "EN translation — Explanation",
     "The English explanation must accurately convey what the JA pattern means and how it is used.",
     "English translation accuracy (6.3)",
     "1. Read the English explanation. "
     "2. Compare against authoritative reference (Genki / Minna no Nihongo lesson reference shown in cross-refs). "
     "3. Verify accuracy, completeness, and absence of misleading claims.",
     "English explanation is accurate, complete, and aligned with authoritative references."),

    ("TS-08", "EN translation — Example sentences",
     "Every example sentence's English translation must be accurate, natural, and convey the same meaning as the JA.",
     "English translation accuracy (6.3)",
     "1. For each example, compare the JA sentence to the EN translation. "
     "2. Verify: meaning, tense, politeness register, naturalness in EN.",
     "All EN translations are accurate AND natural. No literal-translation awkwardness that loses meaning."),

    ("TS-09", "EN translation — Common mistakes commentary",
     "The English commentary on common mistakes must accurately diagnose WHY the wrong line is wrong.",
     "English translation accuracy (6.3)",
     "1. Read the EN explanation under each wrong/right pair. "
     "2. Verify it correctly diagnoses the learner error AND explains the correct usage.",
     "EN commentary accurately diagnoses the error; the explanation matches the right-line correction."),

    ("TS-10", "Meaning line — JA/EN semantic match",
     "The one-line JA meaning (meaning_ja) and one-line EN meaning (meaning_en) at the top of the page must semantically match.",
     "Pattern/explanation consistency (6.2) + EN translation accuracy (6.3)",
     "1. Read the JA meaning line. "
     "2. Read the EN meaning line. "
     "3. Verify they convey the same concept (allowing for natural language difference, not a literal mirror).",
     "JA meaning and EN meaning describe the same grammar concept. No semantic drift."),
]


# ---------- sheet builders ----------

def build_readme(wb):
    ws = wb.create_sheet("ReadMe", 0)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110

    rows = [
        ("Document",        "CategorizedtestScenarios.xlsx — JLPT N5 categorised test scenarios"),
        ("Generated",       date.today().isoformat()),
        ("Source",          "JLPTSuccess N5 — data/grammar.json (178 patterns)"),
        ("Parallel-of",     "specifications/test-scenarios-by-specialist-perspective.xlsx (authoritative bug tracker remains there)"),
        ("Owner",           "Gaurav Srivastava (gaurav.l.srivastava@accenture.com)"),
        ("",                ""),
        ("How to read this file", ""),
        ("Sheet 1: ReadMe",  "You are here. Describes the file."),
        ("Sheet 2: Grammar", "Defines 10 generic test scenarios (TS-01..TS-10) that apply to EACH of the 178 grammar patterns. Read these scenarios first."),
        ("Sheet 3: Grammar Pattern List", "One row per pattern (178 rows). One result column per scenario (TS-01..TS-10). Record Pass/Fail with test date here."),
        ("Sheet 4: Moji",    "Test scenarios for Moji (writing-system / kana / kanji). Phase 2 — scaffold only."),
        ("Sheet 5: Goi",     "Test scenarios for Goi (vocabulary). Phase 2 — scaffold only."),
        ("Sheet 6: Dokkai",  "Test scenarios for Dokkai (reading comprehension). Phase 2 — scaffold only."),
        ("Sheet 7: Chokai",  "Test scenarios for Chokai (listening comprehension). Phase 2 — scaffold only."),
        ("Sheet 8: Test Papers", "Test scenarios for past papers / mock papers. Phase 2 — scaffold only."),
        ("",                ""),
        ("How to run a test", ""),
        ("Step 1",           "Open the Grammar sheet. Pick a scenario (TS-01..TS-10). Read its 'Test steps' column."),
        ("Step 2",           "Open the Grammar Pattern List sheet. Pick a pattern. Click its URL to open the interactive page."),
        ("Step 3",           "Execute the test steps for that scenario against that pattern."),
        ("Step 4",           "Record the result in the corresponding TS-NN column on the Grammar Pattern List row."),
        ("",                ""),
        ("Test result format", ""),
        ("Pass cell",        "Pass (YYYY-MM-DD)"),
        ("Fail cell",        "Fail (YYYY-MM-DD) — short bug description. Full bug details go in the 'Bug Details' column at the right of the row."),
        ("N/A cell",         "N/A (YYYY-MM-DD) — scenario not applicable to this pattern (e.g. pattern has no audio)."),
        ("Blank cell",       "Not yet tested."),
        ("",                ""),
        ("Reference",        "Mandatory test categories from user spec:"),
        ("6.1",              "Japanese-language grammar correctness (sentences, explanations, wrong/right pairs)."),
        ("6.2",              "Pattern name ↔ explanation ↔ example ↔ audio consistency."),
        ("6.3",              "English translation correctness."),
        ("",                ""),
        ("Phase boundary",   "Phase 1 covers Grammar (Sheet 2 + 3). Phase 2 will define scenarios for Moji / Goi / Dokkai / Chokai / Test Papers, and may add authored Japanese-language explanations to the source data."),
        ("Related to",       "Hindi tab is disabled in Phase 1; will be re-enabled in Phase 2 (locale data retained intact in data/locales/hi.json)."),
    ]
    for r_idx, (label, val) in enumerate(rows, start=1):
        ws.cell(row=r_idx, column=1, value=label).font = Font(bold=True, color=GREEN, size=11)
        ws.cell(row=r_idx, column=2, value=val).alignment = WRAP
    ws.cell(row=1, column=1).font = Font(bold=True, color="FFFFFF", size=12)
    ws.cell(row=1, column=1).fill = HEADER_FILL
    ws.cell(row=1, column=2).font = Font(bold=True, color="FFFFFF", size=12)
    ws.cell(row=1, column=2).fill = HEADER_FILL


def build_grammar_scenarios(wb):
    ws = wb.create_sheet("Grammar")
    headers = ["TS ID", "Scenario name", "Description", "Test category (6.1/6.2/6.3)", "Test steps", "Expected result", "Test Result (cross-pattern aggregate)"]
    ws.append(headers)
    style_header_row(ws)
    for ts in GRAMMAR_SCENARIOS:
        ws.append([*ts, ""])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
            cell.border = BORDER
        ws.row_dimensions[row[0].row].height = 70
    auto_width(ws, [9, 38, 55, 28, 70, 50, 30])
    freeze_top(ws)


def build_pattern_list(wb, patterns):
    ws = wb.create_sheet("Grammar Pattern List")
    fixed = ["#", "Pattern ID", "Pattern Name (JA)", "Meaning (EN)", "Category", "Interactive URL (SPA)"]
    ts_cols = [ts[0] for ts in GRAMMAR_SCENARIOS]
    tail = ["Overall (Pass/Fail/Partial)", "Bug Details / Notes"]
    headers = fixed + ts_cols + tail
    ws.append(headers)
    style_header_row(ws)

    patterns_sorted = sorted(patterns, key=lambda p: (p.get("categoryOrder", 999), p.get("patternOrder", 999), p.get("id", "")))
    for i, p in enumerate(patterns_sorted, start=1):
        spa_url = SPA_URL_TEMPLATE.format(id=p["id"])
        row = [
            i,
            p["id"],
            p.get("pattern", ""),
            p.get("meaning_en", ""),
            p.get("category", ""),
            spa_url,
        ] + [""] * len(ts_cols) + ["", ""]
        ws.append(row)
        # hyperlink on URL cell
        url_cell = ws.cell(row=i + 1, column=6)
        url_cell.hyperlink = spa_url
        url_cell.font = Font(color=GREEN_SOFT, underline="single", size=10)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
            cell.border = BORDER

    widths = [5, 11, 22, 38, 32, 60] + [14] * len(ts_cols) + [22, 40]
    auto_width(ws, widths)
    freeze_top(ws)
    ws.freeze_panes = "C2"  # keep row 1 + col A/B sticky for navigation


def build_phase2_scaffold(wb, sheet_name, blurb):
    ws = wb.create_sheet(sheet_name)
    headers = ["TS ID", "Scenario name", "Description", "Test category", "Test steps", "Expected result", "Test Result"]
    ws.append(headers)
    style_header_row(ws)

    note_row = ws.row_dimensions[2]
    note_row.height = 60
    ws.cell(row=2, column=1, value=f"[Phase 2 — scenarios for {sheet_name} to be defined]").font = Font(italic=True, color=GREEN_SOFT, size=11)
    ws.cell(row=2, column=2, value=blurb).alignment = WRAP
    ws.cell(row=2, column=2).font = Font(italic=True, color="555555", size=10)
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=7)
    auto_width(ws, [10, 32, 50, 22, 60, 50, 30])
    freeze_top(ws)


def main():
    with GRAMMAR_JSON.open("r", encoding="utf-8") as f:
        data = json.load(f)
    patterns = data["patterns"]

    wb = Workbook()
    # remove the default sheet then add ours in order
    default_ws = wb.active
    wb.remove(default_ws)

    build_readme(wb)
    build_grammar_scenarios(wb)
    build_pattern_list(wb, patterns)
    build_phase2_scaffold(wb, "Moji",
                          "Moji = writing system (kana, kanji). Future scenarios should check: stroke order, hiragana/katakana correctness, kanji reading accuracy, reading-by-context.")
    build_phase2_scaffold(wb, "Goi",
                          "Goi = vocabulary. Future scenarios should check: word meaning accuracy, example-sentence fit, pitch-accent correctness, vocab_id cross-references resolve.")
    build_phase2_scaffold(wb, "Dokkai",
                          "Dokkai = reading comprehension. Future scenarios should check: passage authenticity, question/answer mapping, distractor plausibility, vocabulary-level appropriateness (JLPT N5).")
    build_phase2_scaffold(wb, "Chokai",
                          "Chokai = listening comprehension. Future scenarios should check: audio clarity, script-audio match, question audio matches written question, distractor plausibility.")
    build_phase2_scaffold(wb, "Test Papers",
                          "Test Papers = past papers / mock papers. Future scenarios should check: paper structure matches official JLPT, answer keys correct, time-limit realistic, score-bands accurate.")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_FILE)
    print(f"OK wrote {OUT_FILE}")
    print(f"   sheets: {wb.sheetnames}")
    print(f"   pattern-list rows: {len(patterns)}   grammar scenarios: {len(GRAMMAR_SCENARIOS)}")


if __name__ == "__main__":
    main()
