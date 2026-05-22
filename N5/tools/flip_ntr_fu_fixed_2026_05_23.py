"""Flip NTR-FU-001 / NTR-FU-002 / NTR-FU-003 (BUG-174/175/176) to
Fixed in the xlsx bug tracker.

Fix Commit cells are set to `<this commit, see git log>` sentinel
shape (JA-146-compliant) — actual hash will be back-filled after
the commit lands.
"""
import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
import openpyxl

REPO_N5 = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(REPO_N5, "specifications", "test-scenarios-by-specialist-perspective.xlsx")

FLIPS = {
    177: ("NTR-FU-001 — Added 11 missing entries to n5_vocab_whitelist.json + new CI invariant JA-151 locks vocab→whitelist reverse direction. README still documents the 4 forward-direction Known mismatches (倍/国籍/週末/では). Asymmetry now fully gated.", "<this commit, see git log>"),
    178: ("NTR-FU-002 — Added 3-band pacing classification: pacing_band_strict (220-240 JEES exam-grade) + pacing_band_ideal (200-220 round-9 ideal) as per-item secondary fields. Distribution: strict in=12/below=38; ideal in=35/below=3/above=12. Keep pacing_status (180-240 learner band, all 50 in_range) for backward compat. _meta.pacing_audit.methodology_three_band_2026_05_23 documents the decision explicitly. Audio re-render at JEES-strict deferred as out-of-scope.", "<this commit, see git log>"),
    179: ("NTR-FU-003 — Renamed collocations → particle_examples corpus-wide on 983 entries (12 pronouns NTR-011 already renamed; total 995/995 unified). Updated js/learn-vocab.js + min.js + en/hi locale strings + CSS class. New CI invariant JA-152 forbids the legacy field name. Fixes silent UI regression from NTR-011 (12 pronouns had lost their particle examples in UI). collocations_provenance preserved as particle_examples_provenance.", "<this commit, see git log>"),
}


def main():
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["User Reported Bugs"]
    for row, (note, commit) in FLIPS.items():
        ws.cell(row, 8).value = "Fixed"
        ws.cell(row, 9).value = commit
        ws.cell(row, 10).value = "2026-05-23"
        # Append fix-note to description
        desc = ws.cell(row, 5).value or ""
        if "**Fix applied" not in desc:
            ws.cell(row, 5).value = desc + f"\n\n**Fix applied 2026-05-23:** {note}"
        print(f"  Flipped row {row} (BUG-{row-3:03d}) to Fixed")
    wb.save(XLSX)
    print()
    print(f"=== Flipped {len(FLIPS)} bugs to Fixed (BUG-174/175/176) ===")
    print("Fix Commit sentinel '<this commit, see git log>' — back-fill after commit lands.")


if __name__ == "__main__":
    main()
