"""Register one BUG row for native-review Batch F (section-level fixes)."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from copy import copy
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / 'specifications' / 'test-scenarios-by-specialist-perspective.xlsx'
wb = load_workbook(XLSX)
ws = wb['User Reported Bugs']
src_row = ws.max_row
new_row = ws.max_row + 1

values = {
    1: '="BUG-"&TEXT(ROW()-3,"000")',
    2: '2026-06-04',
    3: 'Native Japanese / JLPT expert reviewer',
    4: 'Native-review batch F (27 entries): legacy markers + sense disambiguation + pronoun cautions + greeting context + key contrasts + adj-predicate grammar',
    5: (
        'Follow-up to BUG-262 (batches A-D). Batch F targeted the section-'
        'level recommendations in sections 11-30 of the 2026-06-04 native-'
        'reviewer report that did NOT require a second native pass on '
        'replacements (i.e., we add explanatory pragmatic_functions notes, '
        'not new content to be re-verified). 27 entries received '
        'pragmatic_functions additions across 8 sub-batches: '
        '(F1) legacy markers — きしゃ/じびき/せびろ flagged as use-modern-'
        'alternative; (F2) なく — disambiguate 鳴く (animal sound) vs '
        '泣く (human crying) since they share kana but different kanji; '
        '(F3) pronoun usage cautions — かのじょ/かれ (primary modern sense '
        'is girlfriend/boyfriend, not generic she/he), じぶん (most '
        'natural in 自分で / 自分の), あなた (rarely used; prefer name+さん); '
        '(F4) 先生 — honorific scope (teachers, doctors, masters, certain '
        'professionals — not only teacher); (F5) greeting context — '
        'しつれいします (BEFORE), しつれいしました (AFTER), さようなら (more '
        'formal than English goodbye), おかえりなさい (response to ただいま, '
        'home context), おかげさまで (humility frame), ごちそうさまでした '
        '(close-of-meal pair to いただきます); (F6) key contrasts — うみ '
        'vs みずうみ (sea vs lake), ごはん vs こめ (cooked rice/meal vs '
        'uncooked grain), おさけ (alcohol generally, not only Japanese '
        'sake), カップ vs コップ (handled hot-drink cup vs handle-less '
        'tumbler); (F7) な-adj predicate grammar — すき/だいすき/きらい/'
        'だいきらい explained as な-adjectives with が-marked subject, not '
        'English-style verbs; (F8) きれい dual meaning — pretty/beautiful '
        'AND clean/tidy in the same word.'
    ),
    6: 'Medium',
    7: 'P2',
    8: 'Fixed',
    9: '2026-06-04',
    10: '<pending-commit-hash>',
    11: '',  # No new JA invariant — uses existing pragmatic_functions schema
    12: '',
    13: 'BUG-262 (batches A-D), BUG-263 (batch E open)',
    14: 'tools/apply_native_review_batch_f_2026_06_04.py + python tools/check_content_integrity.py (179/179 PASS)',
}

for col in range(1, ws.max_column + 1):
    src_cell = ws.cell(src_row, col)
    new_cell = ws.cell(new_row, col)
    if col in values:
        new_cell.value = values[col]
    if src_cell.has_style:
        new_cell.font = copy(src_cell.font)
        new_cell.alignment = copy(src_cell.alignment)
        if src_cell.fill.fill_type:
            new_cell.fill = copy(src_cell.fill)

wb.save(XLSX)
print(f'Inserted row {new_row}')
