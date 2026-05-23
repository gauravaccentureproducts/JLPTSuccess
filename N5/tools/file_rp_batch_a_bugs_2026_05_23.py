"""File RP-001..005 (Batch A) bugs in the xlsx tracker.

These cover the 5 items closed in Batch A:
  RP-001: goi-4.13 tautological paraphrase
  RP-002: goi-4.6 hospital-worker → doctor inference
  RP-003: 2 dokkai rationale_hi llm_curated → audit-queue
  RP-004: 4 mixed-register rationales
  RP-005: 5 Hindi-punctuation hits

Rejected items (documented in AUDIT-COVERAGE Part 48, not filed as bugs):
  Item 3 — 員 in N5 whitelist (project scope explicitly includes)
  Item 4 — なながつ is defensible distractor
  Item 8/12/15 — auto_inferred is documented provenance, not unverified
  Item 10 — reviewer misread the rationale
  Item 17 — listening.json exists (50 items)
  Item 16 — edge-case flag, not actionable
"""
import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
import openpyxl

REPO_N5 = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(REPO_N5, "specifications", "test-scenarios-by-specialist-perspective.xlsx")

BUGS = [
    {
        "row": 184,
        "title": "RP-001 — goi-4.13 paraphrase tautological (drops よる only)",
        "description": (
            "**Source:** 2026-05-23 reviewer's consolidated 19-item review of v1.16.4.\n\n"
            "**Issue:** goi-4.13 stem 「きのうの よる、はやく ねました」 ⇄ "
            "correct answer 「きのう はやく ねました」. Difference is only "
            "dropping よる. No conceptual jump for the learner. Verified via "
            "paper-4.json inspection.\n\n"
            "**Fix:** rewrote stem to test 「ベッドに 入る」 ≡ 「ねる」 equivalence."
        ),
        "severity": "Medium",
        "priority": "P3",
    },
    {
        "row": 185,
        "title": "RP-002 — goi-4.6 hospital-worker → doctor inference too loose",
        "description": (
            "**Source:** 2026-05-23 review, item 7.\n\n"
            "**Issue:** Stem 「父は びょういんで はたらいて います」 → correct "
            "answer 「父は いしゃです」. Hospital worker isn't necessarily a "
            "doctor (could be nurse / admin / etc.). Inference too loose.\n\n"
            "**Fix:** tightened stem to 「ちちは びょういんで びょうきの 人を "
            "みて います」 (sees sick people at hospital) — makes the inference "
            "to doctor defensible."
        ),
        "severity": "Medium",
        "priority": "P3",
    },
    {
        "row": 186,
        "title": "RP-003 — 2 dokkai rationale_hi entries llm_curated; queue for native review",
        "description": (
            "**Source:** 2026-05-23 review, item 14.\n\n"
            "**Issue:** dokkai-2.3 and dokkai-7.6 carry "
            "`rationale_hi_provenance: 'llm_curated'`. Not native-reviewed; "
            "deferred-by-design per F.44.21 native-speaker audit-block "
            "discipline.\n\n"
            "**Fix:** added `audit.verifier_pending: true` block to both "
            "entries with structured result_schema for the future human "
            "native Hindi/Japanese reviewer."
        ),
        "severity": "Low",
        "priority": "P4",
    },
    {
        "row": 187,
        "title": "RP-004 — 4 mixed-register rationales (English fragments inside Japanese)",
        "description": (
            "**Source:** 2026-05-23 review, item 18.\n\n"
            "**Issue:** Pattern like 「time + に for time of action」 (English "
            "+ Japanese particle mixed). 4 hits: goi-2.4, goi-7.4, goi-7.10, "
            "bunpou-7.1.\n\n"
            "**Fix:** rewrote all 4 rationales as all-Japanese."
        ),
        "severity": "Low",
        "priority": "P4",
    },
    {
        "row": 188,
        "title": "RP-005 — 5 Hindi-punctuation hits (danda after romanized inside parens)",
        "description": (
            "**Source:** 2026-05-23 review, item 19.\n\n"
            "**Issue:** Pattern like 「父 (father)।」 — Hindi danda 「।」 placed "
            "after a parenthetical English gloss inside a Hindi rationale. "
            "5 hits: moji-4.11, moji-5.3, goi-1.4, bunpou-1.10, bunpou-2.4.\n\n"
            "**Fix:** translated the English parenthetical glosses to Hindi."
        ),
        "severity": "Low",
        "priority": "P5",
    },
]


def main():
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["User Reported Bugs"]
    for b in BUGS:
        r = b["row"]
        ws.cell(r, 1).value = '="BUG-"&TEXT(ROW()-3,"000")'
        ws.cell(r, 2).value = "2026-05-23"
        ws.cell(r, 3).value = "reviewer's 19-item consolidated review of v1.16.4"
        ws.cell(r, 4).value = b["title"]
        ws.cell(r, 5).value = b["description"]
        ws.cell(r, 6).value = b["severity"]
        ws.cell(r, 7).value = b["priority"]
        ws.cell(r, 8).value = "Fixed"
        ws.cell(r, 9).value = "2026-05-23"
        ws.cell(r, 10).value = "<this commit, see git log>"
        print(f"  Filed + flipped row {r}: BUG-{r-3:03d} — {b['title'][:60]}")
    wb.save(XLSX)
    print()
    print(f"=== Filed + flipped {len(BUGS)} bugs (BUG-181..185) ===")


if __name__ == "__main__":
    main()
