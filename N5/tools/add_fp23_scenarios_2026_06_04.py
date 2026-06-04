"""Add the xlsx scenario rows JA-116 requires for FP-23 + the new
Phase-0 template-artifact block (batch J doc-propagation)."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from copy import copy
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / 'specifications' / 'test-scenarios-by-specialist-perspective.xlsx'
wb = load_workbook(XLSX)
ws = wb['A. Japanese language']
src_row = ws.max_row  # style template


def write_row(vals):
    new_row = ws.max_row + 1
    for col in range(1, ws.max_column + 1):
        sc = ws.cell(src_row, col)
        nc = ws.cell(new_row, col)
        if col in vals:
            nc.value = vals[col]
        if sc.has_style:
            nc.font = copy(sc.font)
            nc.alignment = copy(sc.alignment)
            if sc.fill.fill_type:
                nc.fill = copy(sc.fill)
    return new_row


r1 = write_row({
    1: 'A-155',
    2: 'False-positive class',
    3: '1',
    4: 'FP-23 template-generated content artifacts — malformed conjugation (<dict-form>ます/ました) + semantic-frame nonsense (今日は とても <adj> / えいがは <adj>でした / この <noun>は <adj>)',
    5: '1. Reference: prompts/Japanese language Accuracy check.txt FP-23. 2. For any template-seeded field (particle_examples, examples), assert no verb entry has form+ます or form+ました (JA-178), and grep example frames for non-fitting headwords. 3. Fix conjugations deterministically; allowlist-remove semantic-frame nonsense; defer naturalness to BUG-263.',
    6: '0 form+ます/form+ました in any verb particle_examples (JA-178 green); 0 えいがは <i-adj>でした; non-weather 今日は とても <adj> removed — against current corpus snapshot.',
    7: 'P1',
    8: 'High',
    9: 'Documentation + audit + CI invariant',
    10: 'BUG-266 batch J; reproduction confirmed 12 cited + 110-verb systematic scope. JA-178 locks the conjugation class. AUDIT-COVERAGE-2026-06-04 Part 64.',
    11: '1.5h',
    12: 'Engineer + Auditor',
    13: 'tools/apply_native_review_batch_j_2026_06_04.py, tools/check_content_integrity.py JA-178',
})

r2 = write_row({
    1: 'Phase-0-template-artifact',
    2: 'Phase-0 block',
    3: '1',
    4: 'Phase-0 template-generated content-artifact block — audit the whole template-seeded field, not just cited instances; fix malformed conjugations deterministically + CI-lock (JA-178); allowlist-remove semantic-frame nonsense',
    5: '1. Reference: prompts/N5Improvement.txt Phase-0 template-generated content-artifact block. 2. Run the JA-178 form+ます/ました check + the example-frame greps. 3. Confirm 0 against current corpus.',
    6: 'JA-178 green (0 malformed verb conjugations); 0 えいがは <i-adj>でした; non-weather 今日は frame removed; カタカナ purchase-template removed.',
    7: 'P1',
    8: 'High',
    9: 'Process + audit',
    10: 'Codifies BUG-266 batch J; companion to F.52 in procedure manual.',
    11: '1h',
    12: 'Engineer + Auditor',
    13: 'tools/check_content_integrity.py, tools/apply_native_review_batch_j_2026_06_04.py',
})

wb.save(XLSX)
print(f'Inserted A-155 (FP-23) row {r1} + Phase-0-template-artifact row {r2}')
