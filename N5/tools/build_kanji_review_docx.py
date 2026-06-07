#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Build a Word (.docx) of N5 kanji entries for native-speaker review.

Mirrors what the SPA kanji page shows in the DEFAULT (English) locale,
one kanji per page, EXCEPT:
  - stroke-order animation / SVG is dropped (stroke COUNT is kept)
  - chrome (nav, back-link, mark-known toggle) is dropped
  - internal metadata (ids, provenance, review_status, reading-passage
    links, recognition_priority, okurigana cut data) is dropped — the
    reviewer checks learner-facing Japanese content, not the plumbing.

Companion to tools/build_vocab_review_docx.py and
tools/build_grammar_review_docx.py — same visual language, same
reviewer-notes box per entry, same TOC bookmark-link scheme, and the
same auto-stamp/keep-exactly-one-dated-copy behaviour (procedure manual
Appendix F.53).

Usage:
  python tools/build_kanji_review_docx.py 一       # one entry by glyph (sample)
  python tools/build_kanji_review_docx.py all      # all N5 kanji
Output: docs/kanji-review/<name>.docx
"""
import json
import sys
import os

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
KANJI = os.path.join(ROOT, "data", "kanji.json")
OUT_DIR = os.path.join(ROOT, "docs", "kanji-review")

CJK = "Yu Gothic"
LATIN = "Calibri"
WRONG_RED = RGBColor(0xB3, 0x1B, 0x1B)
RIGHT_GREEN = RGBColor(0x1A, 0x5C, 0x38)
MUTED = RGBColor(0x66, 0x66, 0x66)
HEAD_DARK = RGBColor(0x14, 0x45, 0x2A)


# --------------------------------------------------------------------------
# Rendering primitives (shared visual language with the vocab/grammar docs)
# --------------------------------------------------------------------------
def _set_fonts(run, cjk=CJK, latin=LATIN):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:ascii'), latin)
    rfonts.set(qn('w:hAnsi'), latin)
    rfonts.set(qn('w:eastAsia'), cjk)


def run(par, text, *, bold=False, italic=False, strike=False, color=None, size=None):
    r = par.add_run(text if text is not None else '')
    r.bold = bold
    r.italic = italic
    if strike:
        r.font.strike = True
    if color is not None:
        r.font.color.rgb = color
    if size is not None:
        r.font.size = Pt(size)
    _set_fonts(r)
    return r


def para(doc, *, space_before=2, space_after=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def subtitle(doc, text):
    if not text:
        return
    p = para(doc, space_after=8)
    run(p, text, italic=True, size=12, color=MUTED)


def section_header(doc, en_label, ja_chip):
    p = para(doc, space_before=10, space_after=3)
    run(p, en_label, bold=True, size=13, color=HEAD_DARK)
    if ja_chip:
        run(p, "  " + ja_chip, bold=False, size=10, color=MUTED)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), 'CCCCCC')
    pbdr.append(bottom); pPr.append(pbdr)


def labeled(doc, label, text):
    if not text:
        return
    p = para(doc)
    run(p, label + ": ", bold=True)
    run(p, text)


def reviewer_box(doc):
    """Reviewer notes table — pre-populated with a visible '未確認 / not yet
    reviewed' placeholder so unreviewed entries flag themselves rather than
    being silently empty. The native reviewer DELETES the placeholder when
    writing 「問題なし。」 or 「要修正: ...」 so the unreviewed queue stays
    visually obvious. (Same convention as the vocab/grammar packets.)"""
    p = para(doc, space_before=10, space_after=2)
    run(p, "Reviewer notes", bold=True, size=11, color=MUTED)
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = 'Table Grid'
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.5)
    cp1 = cell.paragraphs[0]
    cp1.paragraph_format.space_after = Pt(2)
    run(cp1, "☐ 未確認 / not yet reviewed", italic=True, color=WRONG_RED, size=10)
    cp2 = cell.add_paragraph()
    cp2.paragraph_format.space_after = Pt(6)
    run(cp2,
        "[Native reviewer: delete the line above and write 「問題なし。」 if the "
        "entry is correct, OR 「要修正: <issue>」 with a specific issue. Empty "
        "boxes are treated as 'not yet reviewed'.]",
        italic=True, color=MUTED, size=9)
    for _ in range(2):
        cp = cell.add_paragraph()
        cp.paragraph_format.space_after = Pt(6)


# --- clickable Table of Contents: bookmarks + internal hyperlinks ---
_BM_ID = [2000]


def _bookmark(paragraph, name):
    bid = str(_BM_ID[0]); _BM_ID[0] += 1
    start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'), bid); start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), bid)
    pPr = paragraph._p.find(qn('w:pPr'))
    if pPr is not None:
        pPr.addnext(start)
    else:
        paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _toc_link(paragraph, anchor, text):
    h = OxmlElement('w:hyperlink'); h.set(qn('w:anchor'), anchor)
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts')
    rf.set(qn('w:ascii'), LATIN); rf.set(qn('w:hAnsi'), LATIN); rf.set(qn('w:eastAsia'), CJK)
    rPr.append(rf)
    col = OxmlElement('w:color'); col.set(qn('w:val'), '0563C1'); rPr.append(col)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '28'); rPr.append(sz)
    r.append(rPr)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
    r.append(t); h.append(r)
    paragraph._p.append(h)


def ordered_entries(entries):
    """Teaching order: lesson_order ascending, then frequency_rank, then glyph."""
    def key(e):
        return (e.get('lesson_order', 9999), e.get('frequency_rank', 9999), e.get('glyph', ''))
    return sorted(entries, key=key)


def _join_jp(items):
    items = [str(x) for x in (items or []) if x]
    return "、".join(items)


# --------------------------------------------------------------------------
def build_readme(doc, entries, version_meta):
    p = para(doc, space_before=0, space_after=2)
    run(p, "JLPT N5 Kanji — Native-speaker Review Packet", bold=True, size=20, color=HEAD_DARK)
    sv = version_meta.get('site_version', '')
    kv = version_meta.get('kanji_version', '')
    stamp = "  |  ".join([x for x in [
        ("site " + sv) if sv else "",
        ("kanji data " + kv) if kv else "",
        "%d kanji in this packet" % len(entries),
    ] if x])
    if stamp:
        run(para(doc, space_after=8), stamp, italic=True, size=10, color=MUTED)

    section_header(doc, "What this is", "このパケットについて")
    run(para(doc),
        "JLPTSuccess is a free study site for the JLPT N5 (the entry level of the "
        "Japanese-Language Proficiency Test). This document contains every N5 kanji "
        "the site teaches, one per page, exactly as a learner sees it in the default "
        "(English) interface — minus the on-screen chrome and the animated stroke order. "
        "Please review the Japanese for correctness and naturalness.")

    section_header(doc, "What to check on each kanji", "確認のポイント")
    for line in [
        "Readings — are the on'yomi (音) and kun'yomi (訓) correct and complete for N5?",
        "Meaning — does the English meaning match the kanji?",
        "Example words — is each compound real, correctly written, and is its reading (kana) right for that compound?",
        "Example sentences — natural Japanese? Does the English translation match?",
        "Mnemonic / reading rule — is the explanation accurate (especially any reading claim)?",
        "Look-alikes — are the listed kanji genuinely easy to confuse with this one?",
    ]:
        bp = para(doc, space_before=0, space_after=1)
        run(bp, "• " + line)

    section_header(doc, "How to leave feedback", "フィードバックの書き方")
    run(para(doc),
        "Each kanji has a Reviewer notes box. Delete the red 「未確認 / not yet reviewed」 "
        "line and write 「問題なし。」 if it is correct, or 「要修正: …」 with the specific "
        "issue. Boxes left with the red line are treated as not yet reviewed, so progress "
        "stays visible at a glance.")


def build_toc(doc, entries):
    doc.add_page_break()
    run(para(doc, space_before=0, space_after=4), "Contents", bold=True, size=16, color=HEAD_DARK)
    run(para(doc, space_after=8),
        "Click any kanji to jump to its entry (listed in teaching order).",
        italic=True, size=10, color=MUTED)
    PER_ROW = 12
    for start in range(0, len(entries), PER_ROW):
        rowp = para(doc, space_before=2, space_after=2)
        for e in entries[start:start + PER_ROW]:
            _toc_link(rowp, e['_anchor'], e.get('glyph', ''))
            run(rowp, "    ")


def render_entry(doc, e, first):
    if not first:
        doc.add_page_break()

    # Title: the kanji glyph, large
    tp = para(doc, space_before=0, space_after=0)
    run(tp, e.get('glyph', ''), bold=True, size=48, color=HEAD_DARK)
    _bookmark(tp, e['_anchor'])

    # Context line: stroke count + frequency rank (factual, muted)
    bits = []
    if e.get('stroke_count'):
        bits.append("%d strokes" % e['stroke_count'])
    if e.get('frequency_rank'):
        bits.append("frequency rank %d" % e['frequency_rank'])
    if e.get('lesson_order'):
        bits.append("lesson #%d" % e['lesson_order'])
    if bits:
        run(para(doc, space_before=0, space_after=2), "   ".join(bits), size=10, color=MUTED)

    # Subtitle: meanings
    meanings = e.get('meanings') or []
    subtitle(doc, ", ".join(str(m) for m in meanings))

    # READINGS
    section_header(doc, "Readings", "読み")
    labeled(doc, "On'yomi (音)", _join_jp(e.get('on')) or "—")
    labeled(doc, "Kun'yomi (訓)", _join_jp(e.get('kun')) or "—")
    ar = e.get('additional_readings') or {}
    if _join_jp(ar.get('on')):
        labeled(doc, "Additional on'yomi", _join_jp(ar.get('on')))
    if _join_jp(ar.get('kun')):
        labeled(doc, "Additional kun'yomi", _join_jp(ar.get('kun')))

    # MEANING
    section_header(doc, "Meaning", "意味")
    labeled(doc, "English", ", ".join(str(m) for m in meanings) if meanings else "—")
    if e.get('meanings_hi'):
        mh = e['meanings_hi']
        labeled(doc, "Hindi (meanings_hi)", "、".join(mh) if isinstance(mh, list) else str(mh))

    # RADICAL
    rad = e.get('radical')
    if isinstance(rad, dict) and (rad.get('glyph') or rad.get('name')):
        section_header(doc, "Radical", "部首")
        txt = (rad.get('glyph', '') or '')
        if rad.get('name'):
            txt = (txt + "  —  " + rad['name']).strip(" —")
        run(para(doc), txt)

    # EXAMPLE WORDS (compounds that use this kanji)
    examples = e.get('examples') or []
    if examples:
        section_header(doc, "Example words (%d)" % len(examples), "例")
        for ex in examples:
            pp = para(doc, space_before=0, space_after=1)
            run(pp, "•  ")
            run(pp, ex.get('form', '') or '', bold=True)
            if ex.get('reading'):
                run(pp, "  (" + ex['reading'] + ")", color=MUTED)
            if ex.get('gloss'):
                run(pp, "  —  " + ex['gloss'], italic=True, color=MUTED)

    # EXAMPLE SENTENCES
    sentences = e.get('sentences') or []
    if sentences:
        section_header(doc, "Example sentences (%d)" % len(sentences), "例文")
        for s in sentences:
            pp = para(doc)
            run(pp, s.get('ja', ''))
            if s.get('translation_en'):
                run(pp, "  —  " + s['translation_en'], italic=True, color=MUTED)

    # MNEMONIC (learner-facing memory aid; the reading line carries Japanese)
    mn = e.get('mnemonic')
    if isinstance(mn, dict):
        shown = [(k, mn.get(k)) for k in ('summary', 'visual', 'meaning', 'reading') if mn.get(k)]
        # de-dup identical summary/meaning text
        seen = set(); pairs = []
        for k, v in shown:
            if v in seen:
                continue
            seen.add(v); pairs.append((k, v))
        if pairs:
            section_header(doc, "Mnemonic", "覚え方")
            for k, v in pairs:
                labeled(doc, k.capitalize(), v)

    # READING RULE (when does this kanji take on- vs kun-yomi)
    if e.get('reading_rule'):
        section_header(doc, "Reading rule", "読み方の目安")
        run(para(doc), e['reading_rule'])

    # LOOK-ALIKES
    la = [x for x in (e.get('lookalikes') or []) if x]
    if la:
        section_header(doc, "Easy to confuse with", "紛らわしい字")
        run(para(doc), "  ".join(la))

    reviewer_box(doc)


def _load_version_meta():
    out = {}
    try:
        with open(KANJI, encoding='utf-8') as f:
            k = json.load(f)
        out['kanji_version'] = (k.get('_meta') or {}).get('version', '') if isinstance(k, dict) else ''
    except Exception:
        out['kanji_version'] = ''
    try:
        with open(os.path.join(ROOT, 'data', 'version.json'), encoding='utf-8') as f:
            ver = json.load(f)
        out['site_version'] = ver.get('version', '')
        out['counts'] = ver.get('counts', {})
    except Exception:
        out['site_version'] = ''
        out['counts'] = {}
    return out


def main():
    arg = sys.argv[1] if len(sys.argv) > 1 else '一'
    with open(KANJI, encoding='utf-8') as f:
        k = json.load(f)
    entries = k.get('entries') if isinstance(k, dict) else k
    if not isinstance(entries, list):
        print("could not read kanji entries"); return 1

    version_meta = _load_version_meta()

    if arg == 'all':
        ents = ordered_entries(entries)
        out_name = "N5-kanji-%d-entries-for-native-review.docx" % len(ents)
    else:
        ents = [e for e in entries if e.get('glyph') == arg]
        if not ents:
            print("kanji not found:", arg); return 1
        out_name = "N5-kanji-review-SAMPLE-%s.docx" % arg

    # Stable, valid, unique bookmark anchors (start with a letter; index-based
    # so they never collide and never trip Word's bookmark-name rules).
    for i, e in enumerate(ents):
        e['_anchor'] = "k%d" % (i + 1)

    doc = Document()
    st = doc.styles['Normal']
    st.font.name = LATIN
    st.font.size = Pt(11)
    _rpr = st.element.get_or_add_rPr()
    _rfonts = _rpr.find(qn('w:rFonts'))
    if _rfonts is None:
        _rfonts = OxmlElement('w:rFonts')
        _rpr.append(_rfonts)
    _rfonts.set(qn('w:ascii'), LATIN)
    _rfonts.set(qn('w:hAnsi'), LATIN)
    _rfonts.set(qn('w:eastAsia'), CJK)

    build_readme(doc, ents, version_meta)
    if arg == 'all':
        build_toc(doc, ents)
    for i, e in enumerate(ents):
        render_entry(doc, e, first=(i == 0))

    os.makedirs(OUT_DIR, exist_ok=True)
    out_path = os.path.join(OUT_DIR, out_name)
    doc.save(out_path)
    print("wrote", out_path, "(%d entry/entries)" % len(ents))

    # Full review packet only: keep exactly ONE date-stamped copy in the folder.
    # Remove any older stamped copies first, then emit a fresh one stamped with
    # the current date+time, so a reviewer can never be handed a stale upload.
    # (Procedure manual Appendix F.53: dated-deliverable / stale-copy workflow.)
    if arg == 'all':
        import glob, shutil, datetime
        stem = os.path.splitext(out_path)[0]
        for prev in glob.glob(stem + "_*.docx"):
            try:
                os.remove(prev)
            except OSError as e:
                print("  (could not remove %s - file in use? leaving it: %s)" % (os.path.basename(prev), e))
        stamp = datetime.datetime.now().strftime("%Y-%m-%d_%H%M")
        dated = "%s_%s.docx" % (stem, stamp)
        shutil.copy2(out_path, dated)
        print("wrote dated copy", os.path.basename(dated), "(older dated copies removed)")
    return 0


if __name__ == '__main__':
    sys.exit(main())
